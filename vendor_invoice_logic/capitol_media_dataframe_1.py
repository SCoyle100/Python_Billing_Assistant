# capitol_media_dataframe.py

import sys
import re
import pandas as pd
try:
    import win32com.client
except ImportError:
    win32com = None
from docx import Document as DocxDocument

from utils.openai_json import chat_completion_json


INTRO_TOKENS = ("invoice", "start", "weeks", "#", "/")
DISCOUNT_TOKENS = ("discount", "markup", "commission", "rebate", "credit", "adjustment")
AMOUNT_RE = re.compile(r"[-$]?\d[\d,]*(?:\.\d{2})?")

# Reuse your parse_dollar_amount from earlier or define anew:
def parse_dollar_amount(dollar_str: str) -> float:
    cleaned = re.sub(r'[^\d\.]', '', dollar_str)
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def parse_signed_amount(amount_str: str) -> float:
    cleaned = re.sub(r"[^\d\.\-]", "", str(amount_str or ""))
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def clean_word_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r", " ").replace("\x07", " ").replace("\x0b", " ")
    return re.sub(r"\s+", " ", text).strip()


def is_discount_line(text: str) -> bool:
    lowered = str(text or "").strip().lower()
    return any(token in lowered for token in DISCOUNT_TOKENS)


def looks_like_intro_line(text: str) -> bool:
    lowered = str(text or "").strip().lower()
    if not lowered or is_discount_line(lowered):
        return False
    if any(token in lowered for token in INTRO_TOKENS):
        return True
    if lowered.count(",") >= 3:
        return True
    return False


def paragraph_lines_from_com_cell(cell):
    lines = []
    for paragraph_index in range(1, cell.Range.Paragraphs.Count + 1):
        text = clean_word_text(cell.Range.Paragraphs(paragraph_index).Range.Text)
        if text:
            lines.append(text)
    return lines


def paragraph_lines_from_docx_cell(cell):
    lines = []
    for paragraph in cell.paragraphs:
        text = clean_word_text(paragraph.text)
        if text:
            lines.append(text)
    return lines


def row_texts_from_docx_row(row):
    texts = []
    for cell in row.cells:
        text = " ".join(paragraph_lines_from_docx_cell(cell))
        if text:
            texts.append(text)
    return texts


def find_detail_row_index_docx(table):
    for row_index, row in enumerate(table.rows):
        lowered = {text.lower() for text in row_texts_from_docx_row(row)}
        if "description" in lowered and "amount" in lowered:
            return row_index + 1
    return None


def select_detail_cells_docx(detail_row):
    best_desc_idx = None
    best_desc_score = -1
    best_amount_idx = None
    best_amount_score = -1

    for cell_index, cell in enumerate(detail_row.cells):
        lines = paragraph_lines_from_docx_cell(cell)
        amount_score = sum(1 for line in lines if AMOUNT_RE.search(line))
        desc_score = sum(1 for line in lines if any(char.isalpha() for char in line))

        if amount_score > best_amount_score:
            best_amount_idx = cell_index
            best_amount_score = amount_score

        if desc_score > best_desc_score:
            best_desc_idx = cell_index
            best_desc_score = desc_score

    return best_desc_idx, best_amount_idx


def normalize_market_line(text: str) -> str:
    normalized = clean_word_text(text)
    if not normalized:
        return ""

    normalized = re.sub(
        r"\b(?:discount|markup|commission|rebate|credit|adjustment)\b.*$",
        "",
        normalized,
        flags=re.IGNORECASE,
    )
    normalized = normalized.strip(" -,:;")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def extract_market_from_mixed_line(text: str) -> str:
    normalized = normalize_market_line(text)
    if not normalized:
        return ""

    if any(char.isdigit() for char in normalized):
        return ""

    return normalized


def derive_leading_market_from_intro(intro_lines):
    if not intro_lines:
        return ""

    city_list_line = clean_word_text(intro_lines[-1])
    if city_list_line.count(",") < 3:
        return ""

    comma_parts = [part.strip() for part in city_list_line.split(",") if part.strip()]
    if not comma_parts:
        return ""

    first_market = comma_parts[0]
    last_part = comma_parts[-1]

    if last_part.endswith(first_market):
        return first_market

    words = first_market.split()
    if words and last_part.split()[-len(words):] == words:
        return first_market

    return ""


def build_deterministic_capitol_rows_from_docx_table(docx_table):
    detail_row_index = find_detail_row_index_docx(docx_table)
    if detail_row_index is None or detail_row_index >= len(docx_table.rows):
        return [], []

    detail_row = docx_table.rows[detail_row_index]
    detail_desc_idx, detail_amount_idx = select_detail_cells_docx(detail_row)
    if detail_desc_idx is None or detail_amount_idx is None:
        return [], []

    desc_lines = paragraph_lines_from_docx_cell(detail_row.cells[detail_desc_idx])
    amount_lines = paragraph_lines_from_docx_cell(detail_row.cells[detail_amount_idx])

    intro_lines = []
    remaining_desc_lines = list(desc_lines)
    while remaining_desc_lines and looks_like_intro_line(remaining_desc_lines[0]):
        intro_lines.append(remaining_desc_lines.pop(0))

    market_lines = []
    for line in remaining_desc_lines:
        normalized = extract_market_from_mixed_line(line)
        if normalized:
            market_lines.append(normalized)

    adjusted_amounts = []
    for line in amount_lines:
        amount = parse_signed_amount(line)
        if amount > 0:
            adjusted_amounts.append(amount)
        elif amount < 0 and adjusted_amounts:
            adjusted_amounts[-1] += abs(amount)

    if len(market_lines) + 1 == len(adjusted_amounts):
        leading_market = derive_leading_market_from_intro(intro_lines)
        if leading_market:
            market_lines.insert(0, leading_market)

    if len(market_lines) != len(adjusted_amounts):
        print(
            "DEBUG: Deterministic Capitol parser count mismatch - "
            f"markets={len(market_lines)}, adjusted_amounts={len(adjusted_amounts)}"
        )

    invoice_rows = []
    for market, amount in zip(market_lines, adjusted_amounts):
        if amount > 0:
            invoice_rows.append({"Market": market, "Amount": amount})

    return intro_lines, invoice_rows


def identify_invoice_table_with_openai(table_text):
    payload = chat_completion_json(
        system_prompt=(
            "You review table text extracted from a Word document. "
            "Determine whether the table contains invoiceable market rows with monetary amounts. "
            "Return a JSON object with keys: has_invoice_data (boolean), confidence (High, Medium, or Low)."
        ),
        user_prompt=f"Table text:\n{table_text}",
        max_tokens=400,
    )
    return {
        "has_invoice_data": bool(payload.get("has_invoice_data")),
        "confidence": str(payload.get("confidence", "Low")).title(),
    }


def extract_capitol_media_rows_with_openai(table_text):
    payload = chat_completion_json(
        system_prompt=(
            "Extract Capitol Media invoice rows from the provided table text. "
            "Return a JSON object with one key, 'invoices', containing an array of objects. "
            "Each object must use exactly these keys: Description and Amount. "
            "Only include rows that have a real positive billable amount. "
            "Do not include narrative/header text, campaign summary text, city-list intro text, "
            "discount lines, markup lines, commission lines, rebate lines, credit lines, or any negative amounts. "
            "The table may begin with one or two descriptive lines before the market rows start; those belong to the "
            "intro section and must not be included in invoices. "
            "If a line mixes a market name with discount text, keep only the market name. "
            "If the first positive amount corresponds to the first actual market after an intro section, output just the market name."
        ),
        user_prompt=f"Table text:\n{table_text}",
        max_tokens=2500,
    )
    invoices = payload.get("invoices", [])
    if not isinstance(invoices, list):
        raise ValueError("OpenAI Capitol Media extraction did not return a list.")
    return invoices


def find_invoice_table(doc):
    """
    Use OpenAI to dynamically identify which table contains the invoice data.
    """
    print(f"DEBUG: Searching through {doc.Tables.Count} tables for invoice data")
    
    for table_index in range(1, doc.Tables.Count + 1):  # 1-based indexing
        try:
            table = doc.Tables(table_index)
            table_text = table.Range.Text
            print(f"DEBUG: Checking table {table_index}, text length: {len(table_text)}")
            
            result = identify_invoice_table_with_openai(table_text)
            
            print(
                f"DEBUG: Table {table_index} - Has invoice data: {result['has_invoice_data']}, "
                f"Confidence: {result['confidence']}"
            )
            
            if result["has_invoice_data"] and result["confidence"] in ["High", "Medium"]:
                print(f"DEBUG: Selected table {table_index} for invoice processing")
                return table_index, table, table_text
                
        except Exception as e:
            print(f"DEBUG: Error processing table {table_index}: {e}")
            continue
    
    # Fallback to first table if no table is confidently identified
    print("DEBUG: No table confidently identified, using first table as fallback")
    if doc.Tables.Count > 0:
        table = doc.Tables(1)
        return 1, table, table.Range.Text
    
    return None, None, None


def build_dataframe_from_capitol_media(file_path: str) -> pd.DataFrame:
    """
    Dynamically identifies the correct table and reads invoice data,
    then deterministically maps description paragraphs to amounts.

    :param file_path: Path to the .docx file with Capitol Media's table.
    :return: A pandas DataFrame with columns ['Market', 'Amount'].
    """
    try:
        print(f"DEBUG: Processing Capitol Media file: {file_path}")
        
        # 1. Initialize Word
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False  # Make True for debugging

        # 2. Open the document and find the correct table
        doc = word.Documents.Open(file_path)
        
        # Check if document has any tables
        if doc.Tables.Count == 0:
            print("DEBUG: No tables found in document")
            doc.Close(False)
            word.Quit()
            return pd.DataFrame(columns=['Market', 'Amount'])
        
        print(f"DEBUG: Found {doc.Tables.Count} tables in document")
        
        # Dynamically find the table with invoice data
        table_index, table, table_text = find_invoice_table(doc)
        
        if table_index is None or table is None or table_text is None:
            print("DEBUG: No suitable table found")
            doc.Close(False)
            word.Quit()
            return pd.DataFrame(columns=['Market', 'Amount'])
        
        print(f"DEBUG: Table text length: {len(table_text)}")
        print(f"DEBUG: Table text preview: {table_text[:500]}...")

        # 3. Close doc & Word to free resources
        doc.Close(False)
        word.Quit()

        deterministic_rows = []
        intro_lines = []
        try:
            docx_doc = DocxDocument(file_path)
            if 0 < table_index <= len(docx_doc.tables):
                intro_lines, deterministic_rows = build_deterministic_capitol_rows_from_docx_table(
                    docx_doc.tables[table_index - 1]
                )
                print(f"DEBUG: Deterministic intro lines: {intro_lines}")
                print(f"DEBUG: Deterministic invoice row count: {len(deterministic_rows)}")
            else:
                print(
                    f"DEBUG: Selected table index {table_index} is out of range for python-docx tables "
                    f"({len(docx_doc.tables)} found)"
                )
        except Exception as deterministic_exc:
            print(f"DEBUG: Deterministic Capitol parser failed: {deterministic_exc}")

        invoice_rows = []
        if deterministic_rows:
            print("DEBUG: Using deterministic Capitol parser output")
            invoice_rows = deterministic_rows
        else:
            print("DEBUG: Deterministic parser found no rows, falling back to OpenAI extraction...")
            invoices = extract_capitol_media_rows_with_openai(table_text)
            print(f"DEBUG: OpenAI returned {len(invoices)} invoice rows")

            if invoices:
                print(f"DEBUG: Processing {len(invoices)} invoices from OpenAI")
                for i, inv in enumerate(invoices):
                    print(f"DEBUG: Invoice {i}: {inv}")
                    desc_str = inv.get("Description", "").strip()
                    amt_str = inv.get("Amount", "").strip()
                    amt_val = parse_dollar_amount(amt_str)
                    print(f"DEBUG: Parsed - Description: '{desc_str}', Amount: '{amt_str}' -> {amt_val}")

                    if amt_val > 0:
                        invoice_rows.append({"Market": desc_str, "Amount": amt_val})
                        print(f"DEBUG: Added to invoice_rows: Market='{desc_str}', Amount={amt_val}")
            else:
                print("DEBUG: No invoices found in OpenAI result")

        # 6. Create DataFrame
        print(f"DEBUG: Creating DataFrame with {len(invoice_rows)} rows")
        df = pd.DataFrame(invoice_rows)
        
        # Ensure the DataFrame has the expected columns even if empty
        if df.empty:
            df = pd.DataFrame(columns=['Market', 'Amount'])
        
        print(f"DEBUG: Final DataFrame shape: {df.shape}")
        print(f"DEBUG: Final DataFrame columns: {df.columns.tolist()}")
        print(f"DEBUG: Final DataFrame:\n{df}")
        
        return df
    
    except Exception as e:
        print(f"ERROR: Exception in build_dataframe_from_capitol_media: {e}")
        import traceback
        traceback.print_exc()
        # Return empty DataFrame with correct columns on error
        return pd.DataFrame(columns=['Market', 'Amount'])


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python capitol_media_dataframe.py <path_to_docx>")
        sys.exit(1)

    path = sys.argv[1]
    df_invoices = build_dataframe_from_capitol_media(path)
    print("Extracted Invoices:\n", df_invoices)
