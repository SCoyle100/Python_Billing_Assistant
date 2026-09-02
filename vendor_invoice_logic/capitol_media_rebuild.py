from __future__ import annotations

from copy import deepcopy
import logging
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HEADER_KEYWORDS = {"description", "amount"}


def _parse_amount(value):
    cleaned = re.sub(r"[^\d.\-]", "", str(value or ""))
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _format_currency(amount):
    return f"${amount:,.2f}"


def _split_invoice_rows(invoice_rows):
    """
    Convert raw Capitol invoice rows into deterministic display rows.
    Large amounts are split into PART rows capped at $5,000 each.
    """
    display_rows = []
    running_total = 0.0

    for market, amount in invoice_rows:
        normalized_market = str(market or "").strip()
        normalized_amount = _parse_amount(amount)

        if not normalized_market or normalized_amount <= 0:
            continue

        remaining = normalized_amount
        part_index = 0
        if normalized_amount > 5000.0:
            display_rows.append(([normalized_market], ""))
            while remaining > 0:
                part_index += 1
                part_amount = min(5000.0, remaining)
                remaining -= part_amount
                display_rows.append(([f"    - PART {chr(64 + part_index)}"], _format_currency(part_amount)))
                running_total += part_amount
        else:
            display_rows.append(([normalized_market], _format_currency(normalized_amount)))
            running_total += normalized_amount

    return display_rows, running_total


def _is_discount_line(text):
    lowered = str(text or "").strip().lower()
    return "discount" in lowered or "markup" in lowered


def _looks_like_intro_line(text):
    lowered = str(text or "").strip().lower()
    if not lowered or _is_discount_line(lowered):
        return False

    if any(token in lowered for token in ["invoice", "start", "weeks", "#", "/"]):
        return True

    if lowered.count(",") >= 3:
        return True

    return False


def _row_texts(row):
    texts = []
    for cell in row.cells:
        cell_text = " ".join(
            paragraph.text.strip()
            for paragraph in cell.paragraphs
            if paragraph.text and paragraph.text.strip()
        )
        texts.append(cell_text)
    return texts


def _find_main_table(doc):
    for table in doc.tables:
        for row in table.rows:
            lowered = {text.strip().lower() for text in _row_texts(row) if text.strip()}
            if HEADER_KEYWORDS.issubset(lowered):
                return table

    best_table = None
    best_score = -1
    for table in doc.tables:
        score = sum(bool(re.search(r"\d", text)) for row in table.rows for text in _row_texts(row))
        if score > best_score:
            best_table = table
            best_score = score
    return best_table


def _find_header_row_index(table):
    for index, row in enumerate(table.rows):
        lowered = {text.strip().lower() for text in _row_texts(row) if text.strip()}
        if HEADER_KEYWORDS.issubset(lowered):
            return index
    return None


def _find_total_row_index(table, start_index):
    for index in range(start_index, len(table.rows)):
        row_text = " ".join(_row_texts(table.rows[index])).lower()
        if "total" in row_text:
            return index
    return len(table.rows) - 1


def _paragraph_lines(cell):
    return [
        paragraph.text.strip()
        for paragraph in cell.paragraphs
        if paragraph.text and paragraph.text.strip()
    ]


def _extract_intro_lines(detail_row):
    intro_lines = []
    for line in _paragraph_lines(detail_row.cells[0]):
        if _looks_like_intro_line(line):
            intro_lines.append(line)
            continue
        break
    return intro_lines


def _normalize_intro_lines(intro_lines):
    normalized_lines = list(intro_lines)
    if len(normalized_lines) < 2:
        return normalized_lines

    city_line = normalized_lines[-1].strip()
    comma_parts = [part.strip() for part in city_line.split(",") if part.strip()]
    if len(comma_parts) < 2:
        return normalized_lines

    first_city = comma_parts[0]
    last_part = comma_parts[-1]
    first_words = first_city.split()
    last_words = last_part.split()
    if first_words and last_words[-len(first_words):] == first_words:
        trimmed_last_part = " ".join(last_words[:-len(first_words)]).strip(" ,")
        comma_parts[-1] = trimmed_last_part
        normalized_lines[-1] = ", ".join(part for part in comma_parts if part)

    return normalized_lines


def _looks_like_dated_work_order(text):
    normalized = str(text or "").strip().lower()
    has_date = bool(re.search(r"\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b", normalized))
    has_work_order_context = bool(
        re.search(
            r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
            r"jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|"
            r"dec(?:ember)?|w\s*/?\s*o)\b",
            normalized,
        )
    )
    return has_date and has_work_order_context


def _extract_dated_line_item_labels(detail_row):
    """
    Preserve dated work-order labels from the vendor document while ignoring
    their discount/markup rows. Converted Capitol documents align the trailing
    description paragraphs with the amount paragraphs.
    """
    description_text = "\n".join(_paragraph_lines(detail_row.cells[0]))

    # Adobe's Word conversion can merge a work-order label and its discount
    # text into one paragraph (for example, "July wo 7/20 Discount @ 2.5%").
    # Extract the label itself instead of relying on paragraph/amount alignment.
    month_pattern = (
        r"jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
        r"jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|"
        r"dec(?:ember)?"
    )
    label_pattern = re.compile(
        rf"\b(?:{month_pattern})\s+w\s*/?\s*o\s+\d{{1,2}}/\d{{1,2}}(?:/\d{{2,4}})?\b",
        re.IGNORECASE,
    )
    labels = [match.group(0).strip() for match in label_pattern.finditer(description_text)]
    if labels:
        return labels

    # Fallback for converted documents that keep each label and amount on
    # separate, aligned paragraphs.
    description_lines = _paragraph_lines(detail_row.cells[0])
    amount_lines = _paragraph_lines(detail_row.cells[-1])
    if not amount_lines or len(description_lines) < len(amount_lines):
        return []

    aligned_descriptions = description_lines[-len(amount_lines):]
    return [
        description
        for description, amount_text in zip(aligned_descriptions, amount_lines)
        if _parse_amount(amount_text) > 0 and _looks_like_dated_work_order(description)
    ]


def _capture_paragraph_properties(cell):
    properties = []
    for paragraph in cell.paragraphs:
        text = paragraph.text.strip()
        if text and paragraph._p.pPr is not None:
            properties.append(deepcopy(paragraph._p.pPr))
        elif text:
            properties.append(None)
    return properties


def _normalize_invoice_rows_for_display(invoice_rows):
    """Reduce enriched invoice tuples to the market/amount fields used by the table."""
    normalized_rows = []
    for invoice_row in invoice_rows:
        try:
            values = list(invoice_row)
        except TypeError:
            logging.warning("Skipping invalid Capitol invoice row: %r", invoice_row)
            continue

        if len(values) < 2:
            logging.warning("Skipping incomplete Capitol invoice row: %r", invoice_row)
            continue

        normalized_rows.append((values[0], values[1]))

    return normalized_rows


def _apply_source_line_item_labels(invoice_rows, source_labels):
    if not source_labels:
        return invoice_rows

    if len(source_labels) != len(invoice_rows):
        logging.warning(
            "Capitol dated line-item count (%s) did not match invoice row count (%s); "
            "keeping extracted market names.",
            len(source_labels),
            len(invoice_rows),
        )
        return invoice_rows

    return [
        (source_labels[index], amount)
        for index, (_, amount) in enumerate(invoice_rows)
    ]


def _clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag.endswith("tcPr"):
            continue
        tc.remove(child)


def _add_lines_to_cell(
    cell,
    lines,
    *,
    font_name="Times New Roman",
    font_size=9,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    bold=False,
    paragraph_properties=None,
):
    _clear_cell(cell)

    normalized_lines = list(lines) if lines else [""]
    for line_index, line in enumerate(normalized_lines):
        paragraph = cell.add_paragraph()
        if (
            paragraph_properties
            and line_index < len(paragraph_properties)
            and paragraph_properties[line_index] is not None
        ):
            paragraph._p.insert(0, deepcopy(paragraph_properties[line_index]))
        paragraph.alignment = alignment
        run = paragraph.add_run(str(line))
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.bold = bold


def _populate_invoice_row(row, description_lines, amount_text="", description_paragraph_properties=None):
    _add_lines_to_cell(
        row.cells[0],
        description_lines,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        paragraph_properties=description_paragraph_properties,
    )
    _add_lines_to_cell(
        row.cells[-1],
        [amount_text] if amount_text else [""],
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )


def _remove_row_height(row):
    tr = row._tr
    tr_pr = tr.trPr
    if tr_pr is None:
        return

    for child in list(tr_pr):
        if child.tag == qn("w:trHeight"):
            tr_pr.remove(child)


def _set_row_min_height(row, value=360):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(value))
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)


def _set_cell_border(cell, edge, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    border = tc_borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        tc_borders.append(border)

    border.set(qn("w:val"), value)
    if value != "nil":
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def _style_detail_row_borders(row):
    if len(row.cells) < 2:
        return

    left_cell = row.cells[0]
    right_cell = row.cells[-1]

    for cell in (left_cell, right_cell):
        _set_cell_border(cell, "top", "nil")
        _set_cell_border(cell, "bottom", "nil")

    _set_cell_border(left_cell, "left", "single")
    _set_cell_border(left_cell, "right", "single")
    _set_cell_border(right_cell, "left", "single")
    _set_cell_border(right_cell, "right", "single")


def _style_total_row_borders(row):
    if len(row.cells) < 2:
        return

    left_cell = row.cells[0]
    right_cell = row.cells[-1]

    for cell in (left_cell, right_cell):
        _set_cell_border(cell, "top", "single")
        _set_cell_border(cell, "bottom", "single")

    _set_cell_border(left_cell, "left", "single")
    _set_cell_border(left_cell, "right", "single")
    _set_cell_border(right_cell, "left", "single")
    _set_cell_border(right_cell, "right", "single")


def rebuild_capitol_media_table(docx_path, invoice_rows):
    """
    Preserve the converted Word doc's stacked header area, but regenerate the
    detail and total rows under the Description/Amount header.
    """
    doc = Document(docx_path)
    table = _find_main_table(doc)
    if table is None:
        raise ValueError(f"No Capitol Media invoice table found in {docx_path}")

    header_row_index = _find_header_row_index(table)
    if header_row_index is None:
        raise ValueError(f"Could not locate Description/Amount header row in {docx_path}")

    detail_row_index = header_row_index + 1
    if detail_row_index >= len(table.rows):
        raise ValueError(f"Could not locate detail row after Description/Amount header in {docx_path}")

    total_row_index = _find_total_row_index(table, detail_row_index)

    original_detail_row = table.rows[detail_row_index]
    original_total_row = table.rows[total_row_index]
    preserved_intro_lines = _normalize_intro_lines(_extract_intro_lines(original_detail_row))
    intro_paragraph_properties = _capture_paragraph_properties(original_detail_row.cells[0])
    source_line_item_labels = _extract_dated_line_item_labels(original_detail_row)

    detail_template = deepcopy(original_detail_row._tr)
    total_template = deepcopy(original_total_row._tr)

    for row_index in range(total_row_index, detail_row_index - 1, -1):
        table._tbl.remove(table.rows[row_index]._tr)

    # Email overrides enrich Capitol rows with description, job number, and
    # invoice-suffix metadata. The visual pricing table only consumes market
    # and amount, while the full tuples continue downstream to the database.
    display_invoice_rows = _normalize_invoice_rows_for_display(invoice_rows)
    display_invoice_rows = _apply_source_line_item_labels(
        display_invoice_rows,
        source_line_item_labels,
    )

    # The email/extraction rows are authoritative. Vendor discount or markup
    # lines are deliberately omitted rather than folded into the positive rows.
    detail_rows, running_total = _split_invoice_rows(display_invoice_rows)
    row_specs = []

    if preserved_intro_lines:
        intro_display_lines = []
        intro_display_properties = []
        for index, line in enumerate(preserved_intro_lines):
            if index == 1:
                intro_display_lines.append("")
                intro_display_properties.append(
                    intro_paragraph_properties[min(index, len(intro_paragraph_properties) - 1)]
                    if intro_paragraph_properties else None
                )
            intro_display_lines.append(line)
            intro_display_properties.append(
                intro_paragraph_properties[min(index, len(intro_paragraph_properties) - 1)]
                if intro_paragraph_properties else None
            )
        row_specs.append((intro_display_lines, "", intro_display_properties))
        row_specs.extend([([""], "", None), ([""], "", None)])

    for description_lines, amount in detail_rows:
        row_specs.append((description_lines, amount, None))

    row_specs.extend([([""], "", None), ([""], "", None), ([""], "", None)])

    row_templates = [deepcopy(detail_template) for _ in row_specs]
    row_templates.append(deepcopy(total_template))

    header_tr = table.rows[header_row_index]._tr
    for template in reversed(row_templates):
        header_tr.addnext(template)

    inserted_rows_start = detail_row_index
    for offset, (description_lines, amount_text, paragraph_properties) in enumerate(row_specs):
        row = table.rows[inserted_rows_start + offset]
        _remove_row_height(row)
        if paragraph_properties:
            _set_row_min_height(row, 640)
        else:
            _set_row_min_height(row, 220 if not any(line.strip() for line in description_lines) and not amount_text else 360)
        _populate_invoice_row(row, description_lines, amount_text, paragraph_properties)
        _style_detail_row_borders(row)

    total_row = table.rows[inserted_rows_start + len(row_specs)]
    _remove_row_height(total_row)
    _set_row_min_height(total_row, 420)
    _add_lines_to_cell(
        total_row.cells[0],
        [""],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _add_lines_to_cell(
        total_row.cells[-1],
        [f"TOTAL: {_format_currency(running_total)}"],
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
    )
    _style_total_row_borders(total_row)

    doc.save(docx_path)
    logging.info("Rebuilt Capitol Media pricing table in %s", docx_path)
