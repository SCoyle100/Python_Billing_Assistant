import tempfile
import unittest
from pathlib import Path

from docx import Document

from vendor_invoice_logic.capitol_media_rebuild import rebuild_capitol_media_table


class TestCapitolMediaRebuild(unittest.TestCase):
    def test_rebuild_accepts_email_enriched_invoice_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "capitol.docx"
            document = Document()
            table = document.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "Description"
            table.cell(0, 1).text = "Amount"
            description_cell = table.cell(1, 0)
            description_cell.text = "Thompson Tractor - Birmingham Zoo Big Machine Day Radio Campaign"
            description_cell.add_paragraph("7/20 Start, 2 Weeks - Birmingham")
            # Adobe conversion merges each dated work-order label and its
            # discount text into the same paragraph in this vendor document.
            description_cell.add_paragraph("July wo 7/20 Discount @ 2.5%")
            description_cell.add_paragraph("August wo 7/27 Discount @ 2.5%")

            amount_cell = table.cell(1, 1)
            amount_cell.text = "$4,317.00"
            amount_cell.add_paragraph("-$107.92")
            amount_cell.add_paragraph("$683.00")
            amount_cell.add_paragraph("-$17.08")
            table.cell(2, 0).text = "TOTAL"
            table.cell(2, 1).text = "$4,875.00"
            document.save(docx_path)

            enriched_rows = [
                (
                    "Birmingham",
                    4317.0,
                    "",
                    "BIRMINGHAM OUTDOOR",
                    "TTC-500",
                    "",
                ),
                (
                    "Birmingham",
                    683.0,
                    "",
                    "BIRMINGHAM OUTDOOR",
                    "TTC-500",
                    "",
                ),
            ]

            rebuild_capitol_media_table(docx_path, enriched_rows)

            rebuilt = Document(docx_path)
            rebuilt_text = "\n".join(
                cell.text
                for row in rebuilt.tables[0].rows
                for cell in row.cells
            )
            self.assertIn("$4,317.00", rebuilt_text)
            self.assertIn("$683.00", rebuilt_text)
            self.assertIn("TOTAL: $5,000.00", rebuilt_text)
            self.assertIn("July wo 7/20", rebuilt_text)
            self.assertIn("August wo 7/27", rebuilt_text)
            self.assertNotIn("-$107.92", rebuilt_text)
            self.assertNotIn("-$17.08", rebuilt_text)


if __name__ == "__main__":
    unittest.main()
